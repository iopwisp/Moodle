"""The report a student hands in: a normal laboratory report, written from recorded results.

The technical audit (verification checks, evidence register, errors) lives in
``reports.py`` and is written next to it as ``<Assignment>_Audit.docx/.pdf``.
This document is the opposite: a title page, the objective, the materials and
software, the course of work told step by step in plain language (each
integration describes its own steps, see :mod:`lab_agent.narrative`), the
student's own answers, a conclusion and appendices.

What it never does: invent results.  Steps that did not finish are described as
not done, with the reason; figures are embedded only when the registered image
still matches its SHA-256; every number in the text comes from the step's
recorded details.
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from .config import get_config
from .evidence import list_evidence, validate_evidence
from .integrations.base import Narrative, StepFacts
from .models import RunState, TaskStatus
from .narrative import detect_language, join, localize_rows, narrate, shorten_paths

TEXT: dict[str, dict[str, Any]] = {
    "ru": {
        "report": "ОТЧЁТ", "lab": "по лабораторной работе", "discipline": "по дисциплине", "topic": "на тему",
        "student": "Студент", "group": "группа", "id": "ID", "teacher": "Преподаватель", "objective": "Цель работы",
        "materials": "Исходные данные", "software": "Используемое программное обеспечение", "progress": "Ход работы",
        "answers": "Ответы на вопросы", "conclusion": "Вывод", "appendix_files": "Приложение А. Файлы результатов",
        "appendix_commands": "Приложение Б. Выполненные команды", "figure": "Рисунок", "table": "Таблица",
        "command": "Команда", "commands": "Команды", "program": "Программа", "version": "Версия", "purpose": "Назначение",
        "file": "Файл", "size": "Размер, байт", "role": "Роль", "rows_cut": "показаны первые {n} строк из {total}",
        "materials_text": ("Для выполнения работы были выданы файлы, перечисленные в таблице {table}. Перед началом работы они "
                           "скопированы в рабочий каталог, для каждого вычислен SHA-256; оригиналы при выполнении работы не "
                           "изменялись."),
        "software_text": "Работа выполнялась в {os}. Использованные программы приведены в таблице {table}.",
        "all_done": "Все задания лабораторной работы выполнены.",
        "some_open": "Не завершено: {items}. Причины указаны в соответствующих пунктах.",
        "answers_open": "ответы на вопросы",
        "answers_missing": "Ответы на вопросы этого раздела ещё не подготовлены.",
        "files_text": "Ниже перечислены файлы, полученные в ходе работы, с их контрольными суммами SHA-256.",
        "no_commands": "Внешние программы из командной строки не запускались.",
        "roles": {"assignment": "текст задания", "evidence_image": "образ диска", "checksum": "контрольная сумма",
                  "capture": "дамп трафика", "config": "конфигурация", "script": "скрипт", "project": "файл проекта",
                  "image": "изображение", "suspicious": "файл для анализа", "reference": "справочный материал",
                  "archive": "архив", "document": "документ", "data": "данные"},
        "python": "собственные скрипты анализа и проверки", "os_purpose": "операционная система",
    },
    "en": {
        "report": "LABORATORY REPORT", "lab": "", "discipline": "Course", "topic": "Topic", "student": "Student",
        "group": "group", "id": "ID", "teacher": "Instructor", "objective": "Objective", "materials": "Materials",
        "software": "Software", "progress": "Procedure and results", "answers": "Answers to the questions",
        "conclusion": "Conclusion", "appendix_files": "Appendix A. Result files", "appendix_commands": "Appendix B. Commands",
        "figure": "Figure", "table": "Table", "command": "Command", "commands": "Commands", "program": "Program",
        "version": "Version", "purpose": "Purpose", "file": "File", "size": "Size, bytes", "role": "Role",
        "rows_cut": "first {n} of {total} rows shown",
        "materials_text": ("The files listed in Table {table} were supplied. They were copied into the working folder and "
                           "hashed with SHA-256 before the work started; the originals were not modified."),
        "software_text": "The work was done on {os}. Table {table} lists the programs used.",
        "all_done": "All tasks of the assignment were completed.",
        "some_open": "Not completed: {items}. The reasons are given in the corresponding sections.",
        "answers_open": "the answers to the questions",
        "answers_missing": "The answers for this part have not been written yet.",
        "files_text": "The files produced during the work and their SHA-256 hashes are listed below.",
        "no_commands": "No external command-line programs were run.",
        "roles": {"assignment": "assignment text", "evidence_image": "disk image", "checksum": "checksum", "capture": "capture",
                  "config": "configuration", "script": "script", "project": "project file", "image": "image",
                  "suspicious": "file to analyse", "reference": "reference", "archive": "archive"},
        "python": "own analysis and verification scripts", "os_purpose": "operating system",
    },
}
MAX_TABLE_ROWS = 40
TITLE_RE = re.compile(r"^\s*(лабораторная работа|практическая работа|практическое занятие|laboratory work|lab work|lab|"
                      r"assignment|practical work)\s*(?:№|#|no\.?)?\s*(\d+)?\s*[:.\-–—]?\s*(.*)$", re.IGNORECASE)


@dataclass
class Block:
    """One element of the document; both renderers walk the same list."""

    kind: str  # h1 h2 h3 p note label list table figure code caption
    text: str = ""
    rows: list[list[str]] = field(default_factory=list)
    items: list[str] = field(default_factory=list)
    path: Path | None = None
    ordered: bool = False


@dataclass
class StudentReport:
    language: str
    meta: dict[str, str]
    blocks: list[Block]


# ---------------------------------------------------------------------------- markdown
def markdown_blocks(text: str) -> list[Block]:
    """The subset of Markdown students actually write: headings, paragraphs, lists, tables, code, rules."""
    blocks: list[Block] = []
    lines = text.replace("\r\n", "\n").split("\n")
    paragraph: list[str] = []
    index = 0

    def flush() -> None:
        if paragraph:
            blocks.append(Block("p", " ".join(part.strip() for part in paragraph)))
            paragraph.clear()

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if stripped.startswith("```"):
            flush()
            code: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code.append(lines[index])
                index += 1
            blocks.append(Block("code", "\n".join(code)))
        elif re.match(r"^#{1,6}\s", stripped):
            flush()
            level = len(stripped) - len(stripped.lstrip("#"))
            blocks.append(Block({1: "h1", 2: "h2"}.get(level, "h3"), stripped.lstrip("#").strip()))
        elif re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", stripped):
            flush()
        elif stripped.startswith("|") and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[index + 1]):
            flush()
            rows = [[cell.strip() for cell in stripped.strip("|").split("|")]]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            blocks.append(Block("table", rows=rows))
            continue
        elif re.match(r"^\s*([-*+]|\d+[.)])\s+", line):
            flush()
            ordered = bool(re.match(r"^\s*\d+[.)]", line))
            items: list[str] = []
            while index < len(lines) and (re.match(r"^\s*([-*+]|\d+[.)])\s+", lines[index])
                                          or (items and lines[index].startswith("  ") and lines[index].strip())):
                if re.match(r"^\s*([-*+]|\d+[.)])\s+", lines[index]):
                    items.append(re.sub(r"^\s*([-*+]|\d+[.)])\s+", "", lines[index]).strip())
                else:
                    items[-1] += " " + lines[index].strip()
                index += 1
            blocks.append(Block("list", items=items, ordered=ordered))
            continue
        elif not stripped:
            flush()
        else:
            paragraph.append(line)
        index += 1
    flush()
    return blocks


LONG_TOKEN_RE = re.compile(r"\S{28,}")


def has_long_token(text: str) -> bool:
    """Hashes and paths cannot be hyphenated; justified lines around them turn into rivers of white space."""
    return bool(LONG_TOKEN_RE.search(text or ""))


INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|(?<![*\w])\*[^*\s][^*]*\*(?!\w))")


def inline_segments(text: str) -> list[tuple[str, str]]:
    """``**bold**``, ``*italic*`` and ```code``` -> [(text, style)]."""
    segments: list[tuple[str, str]] = []
    position = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > position:
            segments.append((text[position:match.start()], ""))
        token = match.group(0)
        if token.startswith("**"):
            segments.append((token[2:-2], "b"))
        elif token.startswith("`"):
            segments.append((token[1:-1], "code"))
        else:
            segments.append((token[1:-1], "i"))
        position = match.end()
    if position < len(text):
        segments.append((text[position:], ""))
    return segments


# ---------------------------------------------------------------------------- model
def _json(path: Path, default: Any) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return default


def _valid_image(path: Path) -> bool:
    try:
        from PIL import Image

        with Image.open(path) as image:
            image.verify()
        return True
    except (ImportError, OSError, ValueError):
        return False


def _commands(workspace: Path) -> list[tuple[datetime | None, list[str]]]:
    records: list[tuple[datetime | None, list[str]]] = []
    log = workspace / "logs" / "commands.jsonl"
    if log.is_file():
        for line in log.read_text(encoding="utf-8").splitlines():
            try:
                record = json.loads(line)
                stamp = datetime.fromisoformat(str(record.get("timestamp")))
            except (ValueError, TypeError):
                continue
            records.append((stamp, [str(part) for part in record.get("command", [])]))
    return records


def command_line(argv: list[str], workspace: Path) -> str:
    """A command as the student would type it: no WSL wrapper, no absolute paths, short program name."""
    if argv and Path(argv[0]).name.lower() in {"wsl.exe", "wsl"} and "-e" in argv:
        argv = argv[argv.index("-e") + 1:]
    if argv:
        argv = [Path(argv[0].replace("\\", "/")).name, *argv[1:]]
    return shorten_paths(" ".join(f'"{a}"' if " " in a and not a.startswith("--") else a for a in argv), workspace)


def _registry(config: Any) -> Any:
    try:
        from .integrations.registry import build_registry
        from .tools.screenshot import take_screenshot

        return build_registry(take_screenshot, config)
    except Exception:  # noqa: BLE001 - descriptions fall back to generic text
        return None


def _title_parts(title: str) -> tuple[str, str]:
    match = TITLE_RE.match(title or "")
    if match and match.group(3).strip():
        return match.group(2) or "", match.group(3).strip().rstrip(".")
    return "", (title or "").strip().rstrip(".")


def build_report(workspace: Path, registry: Any = None) -> StudentReport:
    state = RunState.model_validate_json((workspace / "state" / "state.json").read_text(encoding="utf-8"))
    config = get_config()
    settings = config.report
    analysis = _json(workspace / "metadata" / "assignment_analysis.json", {})
    environment = _json(workspace / "metadata" / "environment.json", {})
    manifest = _json(workspace / "metadata" / "input_manifest.json", [])
    language = settings.language if settings.language != "auto" else detect_language(
        analysis.get("title", ""), analysis.get("objective", ""), " ".join(analysis.get("questions", [])[:20]),
        " ".join(analysis.get("requirements", [])[:40]))
    text = TEXT[language]
    registry = registry if registry is not None else _registry(config)
    evidence = list_evidence(workspace)
    intact = {item["id"] for item in validate_evidence(workspace) if item.get("verified")}
    commands = _commands(workspace)
    number, topic = _title_parts(analysis.get("title") or state.plan.objective or state.assignment)
    meta = {"number": number, "topic": topic, "course": settings.course, "student": settings.student_name,
            "group": settings.group, "student_id": settings.student_id, "teacher": settings.instructor,
            "university": settings.university, "department": settings.department, "city": settings.city,
            "year": str(datetime.now(UTC).year), "assignment": state.assignment}

    blocks: list[Block] = []
    counters = {"table": 0, "figure": 0, "section": 0}

    def section(title: str) -> None:
        counters["section"] += 1
        blocks.append(Block("h1", f"{counters['section']}. {title}"))

    def table(caption: str, rows: list[list[str]]) -> None:
        if len(rows) < 2:
            return
        counters["table"] += 1
        total = len(rows) - 1
        if total > MAX_TABLE_ROWS:
            rows = rows[:MAX_TABLE_ROWS + 1]
            caption += f" ({text['rows_cut'].format(n=MAX_TABLE_ROWS, total=total)})"
        blocks.append(Block("caption", f"{text['table']} {counters['table']} — {caption}"))
        blocks.append(Block("table", rows=rows))

    def figure(path: Path, caption: str) -> None:
        counters["figure"] += 1
        blocks.append(Block("figure", path=path))
        blocks.append(Block("caption", f"{text['figure']} {counters['figure']} — {caption}"))

    # 1. Objective
    objective = analysis.get("objective") or state.plan.objective or ""
    if objective and objective.strip().rstrip(".") != (analysis.get("title") or "").strip().rstrip("."):
        section(text["objective"])
        blocks.append(Block("p", objective))

    # 2. Materials
    roles = {Path(str(f.get("path", "")).split("::")[-1]).name: f.get("role", "") for f in analysis.get("files", [])}
    material_rows = [[text["file"], text["size"], text["role"], "SHA-256"]]
    for entry in manifest:
        file_name = Path(str(entry.get("workspace_relative") or entry.get("workspace_copy") or entry.get("source") or "")).name
        role = roles.get(file_name, "")
        material_rows.append([file_name, str(entry.get("size", "")), text["roles"].get(role, role), str(entry.get("sha256", ""))])
    if len(material_rows) > 1:
        section(text["materials"])
        blocks.append(Block("p", text["materials_text"].format(table=counters["table"] + 1)))
        table(text["materials"].lower() if language == "en" else "Исходные файлы", material_rows)

    # Descriptions first: they also say which programs each step really used.
    stories: dict[int, Narrative] = {}
    answers: list[tuple[Any, list[Path]]] = []
    for task in state.plan.steps:
        items = [item for item in evidence if item.step_id == task.id]
        markdown = [workspace / item.path for item in items if item.path.lower().endswith(".md") and item.id in intact]
        if task.action == "core.manual_review" and (markdown or analysis.get("questions")):
            answers.append((task, markdown))
            continue
        facts = StepFacts(task.action, task.title, task.status.value, dict(task.parameters), dict(task.result_details),
                          task.status_reason, list(task.requirement_refs), [item.path for item in items],
                          list(task.report_sections), workspace)
        adapter = capability = None
        if registry is not None and registry.has(task.action):
            adapter, capability = registry.adapter_for(task.action), registry.capability(task.action)
        stories[task.id] = narrate(adapter, facts, language, capability.description if capability else "")

    # 3. Software
    attempted = [t for t in state.plan.steps if t.attempts or t.status != TaskStatus.PENDING]
    apps = environment.get("applications") or {}
    used: dict[str, str] = {}
    for task in attempted:
        info = apps.get(task.tool)
        if isinstance(info, dict) and info.get("available"):
            used[str(info.get("display_name") or task.tool)] = str(info.get("version") or "—")
        story = stories.get(task.id)
        for program, version in story.software if story else []:
            used[program] = version if version not in {"", "—"} else used.get(program, "—")
    if any(t.tool in {"forensics", "core"} for t in attempted):
        used.setdefault("Python", str(environment.get("python") or "—"))
    if used:
        section(text["software"])
        blocks.append(Block("p", text["software_text"].format(os=environment.get("os") or "Windows", table=counters["table"] + 1)))
        rows = [[text["program"], text["version"]]] + [[k, v] for k, v in sorted(used.items(), key=lambda kv: kv[0].lower())]
        table(text["software"].lower() if language == "en" else "Программное обеспечение", rows)

    # 4. Course of work
    section(text["progress"])
    sub = 0
    narratives: list[tuple[Any, Narrative]] = []
    for task in state.plan.steps:
        if task.id not in stories:
            continue
        story = stories[task.id]
        items = [item for item in evidence if item.step_id == task.id]
        narratives.append((task, story))
        sub += 1
        blocks.append(Block("h2", f"{counters['section']}.{sub} {story.heading}"))
        for paragraph in story.paragraphs:
            unfinished = task.status != TaskStatus.COMPLETED and paragraph == story.paragraphs[-1] and not story.outcome_explained
            blocks.append(Block("note" if unfinished else "p", paragraph))
        if task.status == TaskStatus.COMPLETED:
            tables = story.tables if story.tables is not None else [
                {"caption": s.get("title", ""), "rows": localize_rows(list(s["table"]), language)}
                for s in task.report_sections if s.get("table")]
            for entry in tables:
                if entry.get("code"):
                    blocks.append(Block("code", str(entry["code"])))
                elif entry.get("rows"):
                    table(str(entry.get("caption", "")), entry["rows"])
            if story.show_code:
                for s in task.report_sections:
                    if s.get("code"):
                        blocks.append(Block("code", str(s["code"])[:4000]))
        if task.started_at is not None:
            end = task.finished_at or datetime.now(UTC)
            used_commands = list(dict.fromkeys(command_line(argv, workspace) for stamp, argv in commands
                                               if stamp is not None and task.started_at <= stamp <= end))
            if used_commands:
                blocks.append(Block("label", (text["command"] if len(used_commands) == 1 else text["commands"]) + ":"))
                blocks.append(Block("code", "\n".join(used_commands)))
        if story.show_figures:
            pictures = [item for item in items if item.type in {"screenshot", "figure"} and item.id in intact
                        and _valid_image(workspace / item.path)]
            for position, item in enumerate(pictures, start=1):
                caption = story.figure_caption or item.description
                figure(workspace / item.path, caption + (f" ({position})" if story.figure_caption and len(pictures) > 1 else ""))

    # 5. Answers written by the student
    if answers:
        section(text["answers"])
        for task, files in answers:
            if not files:
                blocks.append(Block("note", text["answers_missing"]))
                continue
            for path in files:
                parsed = markdown_blocks(path.read_text(encoding="utf-8", errors="replace"))
                if parsed and parsed[0].kind == "h1":
                    parsed = parsed[1:]
                for block in parsed:
                    if block.kind == "h1":
                        block.kind = "h2"
                    blocks.append(block)  # the student's own tables keep their own captions, if any

    # 6. Conclusion
    section(text["conclusion"])
    quote = ("«", "»") if language == "ru" else ("“", "”")
    names = [f"{quote[0]}{story.heading}{quote[1]}" for task, story in narratives
             if task.required and task.status != TaskStatus.COMPLETED]
    if any(task.required and (task.status != TaskStatus.COMPLETED or not files) for task, files in answers):
        names.append(text["answers_open"])
    blocks.append(Block("p", text["some_open"].format(items=join(names, language)) if names else text["all_done"]))
    findings = list(dict.fromkeys(story.finding for _, story in narratives if story.finding))
    if findings:
        blocks.append(Block("p", " ".join(findings)))

    # Appendices
    result_files: list[list[str]] = []
    seen: set[str] = set()
    for item in evidence:
        if item.type in {"screenshot", "figure"} or item.path in seen or item.path.lower().endswith(".md") or item.id not in intact:
            continue
        seen.add(item.path)
        path = workspace / item.path
        result_files.append([item.path, str(path.stat().st_size) if path.is_file() else "", item.sha256 or ""])
    if result_files:
        blocks.append(Block("h1", text["appendix_files"]))
        blocks.append(Block("p", text["files_text"]))
        table(text["appendix_files"].split(". ", 1)[-1], [[text["file"], text["size"], "SHA-256"], *result_files])
    blocks.append(Block("h1", text["appendix_commands"]))
    lines = list(dict.fromkeys(command_line(argv, workspace) for _, argv in commands))
    blocks.append(Block("code", "\n".join(lines)) if lines else Block("p", text["no_commands"]))
    return StudentReport(language, meta, blocks)


# ---------------------------------------------------------------------------- DOCX
def render_docx(report: StudentReport, output: Path) -> Path:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    text, meta = TEXT[report.language], report.meta
    doc = Document()
    page = doc.sections[0]
    page.page_height, page.page_width = Cm(29.7), Cm(21.0)
    page.top_margin = page.bottom_margin = Cm(2)
    page.left_margin, page.right_margin = Cm(3), Cm(1.5)

    def font(style_or_run: Any, name: str = "Times New Roman", size: float = 14, bold: bool | None = None) -> None:
        target = style_or_run.font
        target.name, target.size = name, Pt(size)
        if bold is not None:
            target.bold = bold
        target.color.rgb = RGBColor(0, 0, 0)
        element = style_or_run.element if hasattr(style_or_run, "element") else style_or_run._element
        properties = element.get_or_add_rPr() if hasattr(element, "get_or_add_rPr") else element.rPr
        if properties is not None:
            fonts = properties.find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                properties.append(fonts)
            for key in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                fonts.set(qn(key), name)

    normal = doc.styles["Normal"]
    font(normal)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    for style_name, size in (("Heading 1", 14), ("Heading 2", 14), ("Heading 3", 14)):
        style = doc.styles[style_name]
        font(style, size=size, bold=True)
        style.font.italic = False
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(12), Pt(6)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Cm(1.25)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for style_name in ("List Bullet", "List Number"):
        font(doc.styles[style_name])

    def paragraph(content: str = "", *, align: Any = WD_ALIGN_PARAGRAPH.JUSTIFY, indent: bool = True, size: float = 14,
                  bold: bool = False, italic: bool = False, style: str | None = None) -> Any:
        para = doc.add_paragraph(style=style)
        para.alignment = align
        para.paragraph_format.first_line_indent = Cm(1.25) if indent else Cm(0)
        for part, kind in inline_segments(content):
            run = para.add_run(part)
            font(run, "Courier New" if kind == "code" else "Times New Roman", size - 1 if kind == "code" else size)
            run.bold = bold or kind == "b"
            run.italic = italic or kind == "i"
        return para

    # title page
    centre = WD_ALIGN_PARAGRAPH.CENTER
    for line in (meta["university"], meta["department"]):
        if line:
            paragraph(line.upper() if line is meta["university"] else line, align=centre, indent=False)
    for _ in range(6 if meta["university"] else 9):
        paragraph(indent=False)
    paragraph(text["report"], align=centre, indent=False, size=16, bold=True)
    if report.language == "ru":
        paragraph(f"{text['lab']}{' №' + meta['number'] if meta['number'] else ''}", align=centre, indent=False)
        if meta["course"]:
            paragraph(f"{text['discipline']} «{meta['course']}»", align=centre, indent=False)
        paragraph(f"{text['topic']}: «{meta['topic']}»", align=centre, indent=False)
    else:
        paragraph(meta["topic"], align=centre, indent=False, bold=True)
        if meta["course"]:
            paragraph(f"{text['discipline']}: {meta['course']}", align=centre, indent=False)
    for _ in range(5):
        paragraph(indent=False)
    right = WD_ALIGN_PARAGRAPH.RIGHT
    student = meta["student"] or "____________________"
    paragraph(f"{text['student']}: {student}" + (f", {text['group']} {meta['group']}" if meta["group"] else ""), align=right, indent=False)
    if meta["student_id"]:
        paragraph(f"{text['id']}: {meta['student_id']}", align=right, indent=False)
    paragraph(f"{text['teacher']}: {meta['teacher'] or '____________________'}", align=right, indent=False)
    for _ in range(6):
        paragraph(indent=False)
    paragraph(" ".join(part for part in (meta["city"], meta["year"]) if part), align=centre, indent=False)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # page numbers in the footer; the title page (a "different first page") has none
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = centre
    field_run = footer.add_run()
    for tag, attrs in (("w:fldChar", {"w:fldCharType": "begin"}), ("w:instrText", None), ("w:fldChar", {"w:fldCharType": "end"})):
        node = OxmlElement(tag)
        if attrs:
            for key, value in attrs.items():
                node.set(qn(key), value)
        else:
            node.set(qn("xml:space"), "preserve")
            node.text = "PAGE"
        field_run._r.append(node)
    font(field_run, size=12)
    doc.sections[0].different_first_page_header_footer = True

    for block in report.blocks:
        if block.kind in {"h1", "h2", "h3"}:
            heading = doc.add_heading(level={"h1": 1, "h2": 2, "h3": 3}[block.kind])
            heading.add_run(block.text)
            for run in heading.runs:
                font(run, bold=True)
        elif block.kind in {"p", "note"}:
            align = WD_ALIGN_PARAGRAPH.LEFT if has_long_token(block.text) else WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph(block.text, align=align, italic=block.kind == "note")
        elif block.kind == "label":
            paragraph(block.text, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False)
        elif block.kind == "caption":
            paragraph(block.text, align=centre, indent=False, size=12)
        elif block.kind == "list":
            for number, item in enumerate(block.items, start=1):
                para = paragraph(f"{number}) {item}" if block.ordered else f"– {item}")
                para.paragraph_format.first_line_indent = Cm(1.25)
        elif block.kind == "code":
            for line in (block.text or "").splitlines() or [""]:
                para = doc.add_paragraph()
                para.paragraph_format.line_spacing = 1.0
                para.paragraph_format.first_line_indent = Cm(0)
                run = para.add_run(line or " ")
                font(run, "Courier New", 9.5)
        elif block.kind == "table" and block.rows:
            columns = max(len(row) for row in block.rows)
            grid = doc.add_table(rows=0, cols=columns)
            grid.style = "Table Grid"
            grid.alignment = WD_TABLE_ALIGNMENT.CENTER
            for index, row in enumerate(block.rows):
                cells = grid.add_row().cells
                for cell, value in zip(cells, row + [""] * (columns - len(row)), strict=False):
                    cell.text = ""
                    para = cell.paragraphs[0]
                    para.paragraph_format.line_spacing = 1.0
                    para.paragraph_format.first_line_indent = Cm(0)
                    for part, kind in inline_segments(str(value)):
                        run = para.add_run(part)
                        mono = kind == "code" or bool(re.fullmatch(r"[0-9a-f]{32,64}", part.strip()))
                        font(run, "Courier New" if mono else "Times New Roman", 8 if mono else 11)
                        run.bold = index == 0 or kind == "b"
            doc.add_paragraph().paragraph_format.line_spacing = 1.0
        elif block.kind == "figure" and block.path is not None:
            try:
                from PIL import Image

                with Image.open(block.path) as image:
                    width, height = image.size
                width_cm = 16.0 if height / max(width, 1) < 1.3 else 16.0 * 1.3 * width / max(height, 1)
                doc.add_picture(str(block.path), width=Cm(min(width_cm, 16.0)))
                doc.paragraphs[-1].alignment = centre
                doc.paragraphs[-1].paragraph_format.first_line_indent = Cm(0)
            except Exception:  # noqa: BLE001 - a broken picture must not stop the report
                paragraph(f"[{block.path.name}]", align=centre, indent=False)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    return output


# ---------------------------------------------------------------------------- PDF
def _pdf_fonts() -> tuple[str, str]:
    from reportlab.lib.fonts import addMapping
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    families = [("C:/Windows/Fonts/times.ttf", "C:/Windows/Fonts/timesbd.ttf", "C:/Windows/Fonts/timesi.ttf", "C:/Windows/Fonts/timesbi.ttf"),
                ("/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf",
                 "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Italic.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSerif-BoldItalic.ttf"),
                ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf", "C:/Windows/Fonts/ariali.ttf", "C:/Windows/Fonts/arialbi.ttf")]
    serif = "Times-Roman"
    for paths in families:
        if all(Path(p).is_file() for p in paths):
            serif = "ReportSerif"
            if serif not in pdfmetrics.getRegisteredFontNames():
                for suffix, path in zip(("", "-Bold", "-Italic", "-BoldItalic"), paths, strict=True):
                    pdfmetrics.registerFont(TTFont(serif + suffix, path))
                for bold, italic, suffix in ((0, 0, ""), (1, 0, "-Bold"), (0, 1, "-Italic"), (1, 1, "-BoldItalic")):
                    addMapping(serif, bold, italic, serif + suffix)
            break
    mono = "Courier"
    for path in ("C:/Windows/Fonts/cour.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf", "C:/Windows/Fonts/consola.ttf"):
        if Path(path).is_file():
            mono = "ReportMono"
            if mono not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(mono, path))
            break
    return serif, mono


def render_pdf(report: StudentReport, output: Path) -> Path:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        Image,
        KeepTogether,
        PageBreak,
        Paragraph,
        Preformatted,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    serif, mono = _pdf_fonts()
    text, meta = TEXT[report.language], report.meta
    body = ParagraphStyle("body", fontName=serif, fontSize=13, leading=19.5, alignment=TA_JUSTIFY, firstLineIndent=1.25 * cm)
    note = ParagraphStyle("note", parent=body, fontName=serif + "-Italic" if serif == "ReportSerif" else "Times-Italic")
    centred = ParagraphStyle("centre", parent=body, alignment=TA_CENTER, firstLineIndent=0)
    right = ParagraphStyle("right", parent=body, alignment=TA_RIGHT, firstLineIndent=0)
    caption = ParagraphStyle("caption", parent=centred, fontSize=11.5, leading=15, spaceBefore=3, spaceAfter=8)
    h1 = ParagraphStyle("h1", parent=body, fontName=serif, fontSize=14, leading=20, spaceBefore=14, spaceAfter=6,
                        alignment=TA_LEFT, keepWithNext=1)
    h2 = ParagraphStyle("h2", parent=h1, fontSize=13, spaceBefore=10)
    cell = ParagraphStyle("cell", fontName=serif, fontSize=9.5, leading=11.5, alignment=TA_LEFT, splitLongWords=1)
    code_style = ParagraphStyle("code", fontName=mono, fontSize=8.5, leading=10.5, leftIndent=0.3 * cm)

    def markup(content: str) -> str:
        out = []
        for part, kind in inline_segments(content):
            safe = escape(part)
            out.append(f"<b>{safe}</b>" if kind == "b" else f"<i>{safe}</i>" if kind == "i"
                       else f'<font face="{mono}">{safe}</font>' if kind == "code" else safe)
        return "".join(out)

    story: list[Any] = []
    for line in (meta["university"], meta["department"]):
        if line:
            story.append(Paragraph(escape(line.upper() if line is meta["university"] else line), centred))
    story.append(Spacer(1, 6 * cm if meta["university"] else 8 * cm))
    story.append(Paragraph(f"<b>{escape(text['report'])}</b>", ParagraphStyle("t", parent=centred, fontSize=16, leading=24)))
    if report.language == "ru":
        story.append(Paragraph(escape(f"{text['lab']}{' №' + meta['number'] if meta['number'] else ''}"), centred))
        if meta["course"]:
            story.append(Paragraph(escape(f"{text['discipline']} «{meta['course']}»"), centred))
        story.append(Paragraph(escape(f"{text['topic']}: «{meta['topic']}»"), centred))
    else:
        story.append(Paragraph(f"<b>{escape(meta['topic'])}</b>", centred))
        if meta["course"]:
            story.append(Paragraph(escape(f"{text['discipline']}: {meta['course']}"), centred))
    story.append(Spacer(1, 3.5 * cm))
    student = meta["student"] or "____________________"
    story.append(Paragraph(escape(f"{text['student']}: {student}" + (f", {text['group']} {meta['group']}" if meta["group"] else "")), right))
    if meta["student_id"]:
        story.append(Paragraph(escape(f"{text['id']}: {meta['student_id']}"), right))
    story.append(Paragraph(escape(f"{text['teacher']}: {meta['teacher'] or '____________________'}"), right))
    story.append(Spacer(1, 4 * cm))
    story.append(Paragraph(escape(" ".join(p for p in (meta["city"], meta["year"]) if p)), centred))
    story.append(PageBreak())

    width = A4[0] - 4.5 * cm
    pending_caption: Any = None
    for block in report.blocks:
        if block.kind == "h1":
            story.append(Paragraph(f"<b>{escape(block.text)}</b>", h1))
        elif block.kind in {"h2", "h3"}:
            story.append(Paragraph(f"<b>{escape(block.text)}</b>", h2))
        elif block.kind in {"p", "note"}:
            style = note if block.kind == "note" else body
            if has_long_token(block.text):
                style = ParagraphStyle(style.name + "-left", parent=style, alignment=TA_LEFT)
            story.append(Paragraph(markup(block.text), style))
        elif block.kind == "label":
            story.append(Paragraph(markup(block.text), ParagraphStyle("label", parent=body, firstLineIndent=0, alignment=TA_LEFT)))
        elif block.kind == "list":
            for number, item in enumerate(block.items, start=1):
                story.append(Paragraph(markup(f"{number}) {item}" if block.ordered else f"– {item}"), body))
        elif block.kind == "code":
            story.append(Spacer(1, 3))
            story.append(Preformatted(block.text or " ", code_style, maxLineLength=95, newLineChars="    "))
            story.append(Spacer(1, 5))
        elif block.kind == "caption":
            if story and isinstance(story[-1], Image):
                image = story.pop()
                story.append(KeepTogether([image, Paragraph(escape(block.text), caption)]))
            else:
                pending_caption = Paragraph(escape(block.text), ParagraphStyle("tc", parent=caption, alignment=TA_LEFT))
        elif block.kind == "table" and block.rows:
            columns = max(len(row) for row in block.rows)
            rows = [row + [""] * (columns - len(row)) for row in block.rows]
            # a column is at least as wide as its longest word (hashes excepted: they may wrap), otherwise
            # proportional to its content, so that "Инструмент" or "00000384.png" are never broken mid-word
            longest = [min(30, max(len(word) for row in rows for word in (str(row[i]).split() or [""]))) for i in range(columns)]
            content = [min(45, max(len(str(row[i])) for row in rows)) for i in range(columns)]
            weights = [max(4.0, longest[i] * 1.2, content[i] * 0.8) for i in range(columns)]
            widths = [width * w / sum(weights) for w in weights]
            data = [[Paragraph(markup(str(value)) if index else f"<b>{escape(str(value))}</b>", cell) for value in row]
                    for index, row in enumerate(rows)]
            grid = Table(data, colWidths=widths, repeatRows=1, hAlign="CENTER")
            grid.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                      ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3)]))
            story += [pending_caption, grid, Spacer(1, 8)] if pending_caption is not None else [grid, Spacer(1, 8)]
            pending_caption = None
        elif block.kind == "figure" and block.path is not None:
            try:
                picture = Image(str(block.path))
                picture._restrictSize(width, 18 * cm)
                story.append(picture)
            except Exception:  # noqa: BLE001
                story.append(Paragraph(escape(f"[{block.path.name}]"), centred))

    def number_pages(canvas: Any, document: Any) -> None:
        canvas.saveState()
        canvas.setFont(serif, 11)
        canvas.drawCentredString(A4[0] / 2, 1.2 * cm, str(document.page))
        canvas.restoreState()

    output.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(output), pagesize=A4, leftMargin=3 * cm, rightMargin=1.5 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
                      title=meta["topic"], author=meta["student"]).build(story, onLaterPages=number_pages)
    return output


def generate_student_report(workspace: Path, registry: Any = None) -> dict[str, str]:
    state = RunState.model_validate_json((workspace / "state" / "state.json").read_text(encoding="utf-8"))
    report = build_report(workspace, registry)
    folder = workspace / "reports"
    docx = render_docx(report, folder / f"{state.assignment}_Report.docx")
    pdf = render_pdf(report, folder / f"{state.assignment}_Report.pdf")
    return {"docx": str(docx), "pdf": str(pdf)}
