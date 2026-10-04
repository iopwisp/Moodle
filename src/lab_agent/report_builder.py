"""Build a student report from one Markdown file, for work done by hand or outside the runner.

Instead of a new ``make_report.py`` per assignment, the report is written as Markdown with a short
YAML header and rendered with the same DOCX/PDF layout as the agent's own student report::

    ---
    title: Анализ и восстановление разделов MBR и GPT
    number: 4
    course: Introduction to Digital Forensics
    language: ru
    output: Assignment4_{student_id}_{surname}.docx
    ---
    # 1. Цель работы
    Text with **bold**, *italic* and `code`.

    ![Окно TestDisk после анализа](screenshots/testdisk.png)

    Таблица: Контрольные суммы образов
    ![](report_data/source_hashes.csv)

Pictures and CSV tables are numbered automatically ("Рисунок 1 — ...", "Таблица 1 — ..."); a missing
or broken picture is an error, never a placeholder.  Student name, group and ID come from the
``report`` section of config.yaml unless the header overrides them.  :func:`lint` flags phrasing that
makes a report read as machine-written, and leftovers such as TODO or "[вставить ...]".
"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import yaml

from .config import get_config
from .student_report import TEXT, Block, StudentReport, markdown_blocks, render_docx, render_pdf

IMAGE_LINE_RE = re.compile(r"^\s*!\[(?P<caption>[^\]]*)\]\((?P<path>[^)]+)\)\s*$")
TABLE_CAPTION_RE = re.compile(r"^\s*(?:Table|Таблица)\s*:\s*(?P<caption>.+?)\s*$", re.IGNORECASE)
META_KEYS = ("title", "number", "course", "student", "surname", "group", "student_id", "teacher", "university",
             "department", "city", "year", "language", "output")
MAX_CSV_ROWS = 60


class ReportSpecError(ValueError):
    """The Markdown report cannot be built as written (missing picture, bad header, ...)."""


@dataclass
class LintFinding:
    level: str  # error | warning
    message: str
    excerpt: str = ""

    def __str__(self) -> str:
        return f"[{self.level}] {self.message}" + (f": «{self.excerpt}»" if self.excerpt else "")


# ---------------------------------------------------------------------------- parsing
def split_front_matter(text: str) -> tuple[dict[str, Any], str]:
    text = text.lstrip("﻿").replace("\r\n", "\n")
    if not text.startswith("---\n"):
        return {}, text
    end = text.find("\n---", 4)
    if end < 0:
        raise ReportSpecError("The YAML header starts with --- but has no closing --- line")
    try:
        header = yaml.safe_load(text[4:end]) or {}
    except yaml.YAMLError as exc:
        raise ReportSpecError(f"The YAML header is not valid: {exc}") from exc
    if not isinstance(header, dict):
        raise ReportSpecError("The YAML header must be a set of 'key: value' lines")
    unknown = sorted(set(header) - set(META_KEYS))
    if unknown:
        raise ReportSpecError(f"Unknown header field(s) {unknown}; allowed: {', '.join(META_KEYS)}")
    body = text[end + 4:]
    body = body.removeprefix("\n")
    return {key: "" if value is None else str(value) for key, value in header.items()}, body


def _isolate_special_lines(body: str) -> str:
    """Put picture and table-caption lines in paragraphs of their own so Markdown does not merge them with text."""
    out: list[str] = []
    fence = False
    for line in body.split("\n"):
        if line.strip().startswith("```"):
            fence = not fence
        if not fence and (IMAGE_LINE_RE.match(line) or TABLE_CAPTION_RE.match(line)):
            out.extend(["", line.strip(), ""])
        else:
            out.append(line)
    return "\n".join(out)


def _csv_rows(path: Path) -> tuple[list[list[str]], int]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        sample = stream.read(4096)
        stream.seek(0)
        try:
            dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = [row for row in csv.reader(stream, dialect) if any(cell.strip() for cell in row)]
    if len(rows) < 2:
        raise ReportSpecError(f"{path.name} has no data rows to show as a table")
    total = len(rows) - 1
    return rows[:MAX_CSV_ROWS + 1], total


def _check_image(path: Path) -> None:
    try:
        from PIL import Image

        with Image.open(path) as image:
            image.verify()
    except ImportError:
        if path.stat().st_size == 0:
            raise ReportSpecError(f"Picture {path} is empty") from None
    except Exception as exc:
        raise ReportSpecError(f"Picture {path} is not a valid image ({exc})") from exc


def build_spec_report(source: Path) -> tuple[StudentReport, dict[str, str]]:
    """Parse a Markdown report into the shared document model; returns (report, header)."""
    header, body = split_front_matter(source.read_text(encoding="utf-8"))
    settings = get_config().report
    language = header.get("language") or (settings.language if settings.language != "auto" else "")
    if not language:
        language = "ru" if re.search(r"[А-Яа-яЁё]", body) else "en"
    if language not in TEXT:
        raise ReportSpecError(f"language must be one of {sorted(TEXT)}, not {language!r}")
    text = TEXT[language]
    meta = {"number": header.get("number", ""), "topic": header.get("title") or source.stem,
            "course": header.get("course", settings.course), "student": header.get("student", settings.student_name),
            "group": header.get("group", settings.group), "student_id": header.get("student_id", settings.student_id),
            "teacher": header.get("teacher", settings.instructor), "university": header.get("university", settings.university),
            "department": header.get("department", settings.department), "city": header.get("city", settings.city),
            "year": header.get("year") or str(datetime.now(UTC).year), "assignment": source.stem}

    blocks: list[Block] = []
    counters = {"figure": 0, "table": 0}
    pending_caption = ""
    for block in markdown_blocks(_isolate_special_lines(body)):
        caption_match = TABLE_CAPTION_RE.match(block.text) if block.kind == "p" else None
        image_match = IMAGE_LINE_RE.match(block.text) if block.kind == "p" else None
        if caption_match:
            pending_caption = caption_match.group("caption")
            continue
        if image_match:
            path = (source.parent / image_match.group("path").strip()).resolve()
            if not path.is_file():
                raise ReportSpecError(f"File not found: {image_match.group('path')} (relative to {source.parent})")
            if path.suffix.lower() == ".csv":
                rows, total = _csv_rows(path)
                caption = pending_caption or image_match.group("caption")
                if total > MAX_CSV_ROWS:
                    caption += f" ({text['rows_cut'].format(n=MAX_CSV_ROWS, total=total)})"
                counters["table"] += 1
                blocks.append(Block("caption", f"{text['table']} {counters['table']} — {caption}".rstrip(" —")))
                blocks.append(Block("table", rows=rows))
            else:
                _check_image(path)
                counters["figure"] += 1
                blocks.append(Block("figure", path=path))
                caption = image_match.group("caption").strip()
                blocks.append(Block("caption", f"{text['figure']} {counters['figure']}" + (f" — {caption}" if caption else "")))
            pending_caption = ""
            continue
        if block.kind == "table" and pending_caption:
            counters["table"] += 1
            blocks.append(Block("caption", f"{text['table']} {counters['table']} — {pending_caption}"))
            pending_caption = ""
        blocks.append(block)
    if pending_caption:
        raise ReportSpecError(f"Table caption «{pending_caption}» is not followed by a table")
    return StudentReport(language, meta, blocks), header


def output_name(header: dict[str, str], report: StudentReport, source: Path) -> str:
    template = header.get("output") or f"{source.stem}.docx"
    meta = report.meta
    student = meta.get("student", "")
    # "Фамилия Имя" is the usual order in config.yaml; set "surname:" in the header when it is not.
    values = {"student": student.replace(" ", "_"), "surname": header.get("surname") or (student.split() or ["Student"])[0],
              "group": meta.get("group", ""), "student_id": meta.get("student_id", "") or "StudentID",
              "title": re.sub(r"\W+", "_", meta.get("topic", ""))[:40].strip("_")}
    try:
        name = template.format(**values)
    except (KeyError, IndexError) as exc:
        raise ReportSpecError(f"output may use {{{'}, {'.join(values)}}}, not {exc}") from exc
    name = re.sub(r'[<>:"/\\|?*]+', "_", name)
    return name if name.lower().endswith(".docx") else name + ".docx"


def build_report_files(source: Path, out_dir: Path | None = None, *, pdf: bool = True) -> dict[str, Any]:
    report, header = build_spec_report(source)
    folder = out_dir or source.parent
    docx = render_docx(report, folder / output_name(header, report, source))
    result: dict[str, Any] = {"docx": str(docx), "figures": sum(b.kind == "figure" for b in report.blocks),
                              "tables": sum(b.kind == "table" for b in report.blocks), "lint": lint_report(report)}
    if pdf:
        result["pdf"] = str(render_pdf(report, docx.with_suffix(".pdf")))
    return result


# ---------------------------------------------------------------------------- style check
# Phrases that make a lab report read as generated: filler openers, evaluative adverbs, stock transitions.
AI_PHRASES: dict[str, tuple[str, ...]] = {
    "ru": ("в данной лабораторной работе", "в рамках данной работы", "в ходе данной работы", "успешно",
           "стоит отметить", "важно отметить", "следует отметить", "необходимо отметить", "таким образом",
           "в заключение", "подводя итог", "играет ключевую роль", "играет важную роль", "является неотъемлемой",
           "позволяет эффективно", "комплексный подход", "всесторонн", "в современном мире",
           "наглядно демонстрирует", "бесшовн", "ключевой аспект", "давайте"),
    "en": ("in this lab", "in this assignment we", "successfully", "it is worth noting", "it is important to note",
           "in conclusion", "to sum up", "overall,", "furthermore", "moreover", "plays a crucial role", "crucial",
           "delve", "comprehensive", "seamless", "robust", "leverage", "in today's", "showcas", "pivotal",
           "a testament to", "let's"),
}
LEFTOVER_RE = re.compile(r"\b(TODO|TBD|FIXME|XXX|lorem ipsum)\b|\[(вставить|insert|добавить|скриншот|screenshot)[^\]]*\]|<[^<>]{0,40}(здесь|here)[^<>]{0,40}>",
                         re.IGNORECASE)


def lint_text(paragraphs: list[str], list_items: int, language: str) -> list[LintFinding]:
    findings: list[LintFinding] = []
    joined = "\n".join(paragraphs)
    for match in LEFTOVER_RE.finditer(joined):
        findings.append(LintFinding("error", "unfinished placeholder left in the text", match.group(0)))
    lowered = joined.casefold()
    for phrase in AI_PHRASES.get(language, ()) + (AI_PHRASES["en"] if language != "en" else ()):
        count = lowered.count(phrase)
        if count:
            start = lowered.find(phrase)
            excerpt = joined[max(0, start - 30): start + len(phrase) + 30].replace("\n", " ").strip()
            findings.append(LintFinding("warning", f"stock phrase “{phrase}” ×{count} — rewrite in your own words", excerpt))
    openers: dict[str, int] = {}
    for paragraph in paragraphs:
        words = re.findall(r"\w+", paragraph.casefold())[:2]
        if len(words) == 2:
            openers[" ".join(words)] = openers.get(" ".join(words), 0) + 1
    for opener, count in openers.items():
        if count >= 3:
            findings.append(LintFinding("warning", f"{count} paragraphs start with the same words", opener))
    if list_items > 12 and list_items > 2 * max(1, len(paragraphs)):
        findings.append(LintFinding("warning", f"{list_items} list items vs {len(paragraphs)} paragraphs — a lab report "
                                               "describes the work in sentences, not only bullet points"))
    return findings


def lint_report(report: StudentReport) -> list[str]:
    paragraphs = [block.text for block in report.blocks if block.kind in {"p", "note"}]
    items = sum(len(block.items) for block in report.blocks if block.kind == "list")
    paragraphs += [item for block in report.blocks if block.kind == "list" for item in block.items]
    return [str(item) for item in lint_text(paragraphs, items, report.language)]


def lint_file(path: Path) -> list[str]:
    """Style check of a Markdown report or a finished .docx file."""
    if path.suffix.lower() == ".docx":
        from docx import Document

        document = Document(str(path))
        paragraphs = [p.text for p in document.paragraphs if len(p.text.split()) > 3]
        items = sum(1 for p in document.paragraphs if p.style is not None and p.style.name.startswith("List"))
        language = "ru" if re.search(r"[А-Яа-яЁё]", " ".join(paragraphs)) else "en"
        return [str(item) for item in lint_text(paragraphs, items, language)]
    report, _ = build_spec_report(path)
    return lint_report(report)
