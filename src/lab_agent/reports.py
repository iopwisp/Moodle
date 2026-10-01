"""Technical audit report (DOCX/PDF) built only from recorded results.

The document a student hands in is written by :mod:`lab_agent.student_report`;
this one is the full record behind it: every step with its capability and
parameters, verification checks, requirement mapping, evidence register with
hashes, errors and recoveries.

Both formats render the same :class:`ReportModel`, assembled from the
checkpoint, the evidence registry, the input manifest, environment discovery
and adapter-provided report sections.  Nothing is invented: steps that were not
verified are reported as FAILED/BLOCKED with their reasons, and screenshots are
embedded only when the registered file still matches its SHA-256.
"""

from __future__ import annotations

import json
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from .config import get_config
from .credentials import redact
from .evidence import list_evidence, validate_evidence
from .models import EvidenceItem, RunState, TaskStatus

LABELS = {
    "en": {
        "report": "Laboratory Report", "student": "Student", "group": "Group", "student_id": "Student ID", "case": "Case ID",
        "course": "Course", "date": "Date", "status": "Final status", "objective": "1. Objective", "environment": "2. Environment",
        "inputs": "3. Input materials and integrity", "method": "4. Method", "execution": "5. Execution and results",
        "mapping": "6. Requirement mapping", "evidence": "7. Evidence register", "errors": "8. Errors and recoveries",
        "conclusion": "9. Conclusion", "facts": "Facts", "limitations": "Limitations", "final": "Final result",
        "appendix": "Appendix A. Commands executed", "action": "Action", "result": "Result", "verification": "Verification",
        "figure": "Figure", "not_set": "(not set - configure report.* in config.yaml)",
    },
    "ru": {
        "report": "Отчёт по лабораторной работе", "student": "Студент", "group": "Группа", "student_id": "Student ID",
        "case": "Идентификатор кейса", "course": "Курс", "date": "Дата", "status": "Итоговый статус", "objective": "1. Цель работы",
        "environment": "2. Среда выполнения", "inputs": "3. Исходные материалы и целостность", "method": "4. Методика",
        "execution": "5. Ход работы и результаты", "mapping": "6. Соответствие требованиям", "evidence": "7. Реестр доказательств",
        "errors": "8. Ошибки и восстановление", "conclusion": "9. Заключение", "facts": "Факты", "limitations": "Ограничения",
        "final": "Итоговый вывод", "appendix": "Приложение A. Выполненные команды", "action": "Действие", "result": "Результат",
        "verification": "Проверка", "figure": "Рисунок", "not_set": "(не указано - заполните report.* в config.yaml)",
    },
}


@dataclass
class Figure:
    path: Path
    caption: str


@dataclass
class StepView:
    id: int
    title: str
    status: str
    capability: str
    parameters: dict[str, Any]
    summary: str
    reason: str
    requirement_refs: list[str]
    checks: list[dict[str, Any]]
    sections: list[dict[str, Any]]
    figures: list[Figure]
    evidence: list[EvidenceItem]
    attempts: int
    required: bool


@dataclass
class ReportModel:
    title: str
    assignment: str
    labels: dict[str, str]
    cover: list[tuple[str, str]]
    status: str
    status_reason: str
    objective: str
    environment: list[list[str]]
    os_line: str
    inputs: list[list[str]]
    method: list[list[str]]
    steps: list[StepView]
    mapping: list[list[str]]
    evidence: list[list[str]]
    errors: list[list[str]]
    recoveries: list[list[str]]
    facts: list[str]
    limitations: list[str]
    final: str
    commands: list[str] = field(default_factory=list)


def _state(workspace: Path) -> RunState:
    return RunState.model_validate_json((workspace / "state" / "state.json").read_text(encoding="utf-8"))


def _output_paths(workspace: Path, state: RunState) -> tuple[Path, Path]:
    folder = workspace / "reports"
    folder.mkdir(parents=True, exist_ok=True)
    return folder / f"{state.assignment}_Audit.docx", folder / f"{state.assignment}_Audit.pdf"


def _valid_image(path: Path) -> bool:
    try:
        from PIL import Image

        with Image.open(path) as image:
            image.verify()
        return True
    except (ImportError, OSError, ValueError):
        return False


def _json(path: Path, default: Any) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return default


def build_model(workspace: Path) -> ReportModel:
    state = _state(workspace)
    config = get_config()
    labels = LABELS.get(config.report.language, LABELS["en"])
    evidence = list_evidence(workspace)
    integrity = {entry["id"]: entry for entry in validate_evidence(workspace)}
    analysis = _json(workspace / "metadata" / "assignment_analysis.json", {})
    environment = _json(workspace / "metadata" / "environment.json", {})
    manifest = _json(workspace / "metadata" / "input_manifest.json", [])
    report = config.report
    not_set = labels["not_set"]
    title = analysis.get("title") or state.plan.objective or state.assignment
    cover = [
        (labels["course"], report.course or not_set), (labels["student"], report.student_name or not_set),
        (labels["group"], report.group or not_set), (labels["student_id"], report.student_id or not_set),
        (labels["case"], report.case_id or f"CASE-{state.assignment}"),
        (labels["date"], datetime.now(UTC).strftime("%Y-%m-%d")), (labels["status"], str(state.status)),
    ]
    used_tools = {task.tool for task in state.plan.steps}
    env_rows = [["Application", "Available", "Version", "Path"]]
    for name, app in sorted((environment.get("applications") or {}).items()):
        if app.get("available") or name in used_tools:
            env_rows.append([app.get("display_name") or name, "yes" if app.get("available") else "no",
                             str(app.get("version") or "-"), str(app.get("path") or app.get("note") or "-")])
    inputs = [["File", "Size", "SHA-256", "Origin"]]
    for entry in manifest:
        inputs.append([Path(str(entry.get("workspace_relative") or entry.get("workspace_copy") or "")).name or str(entry.get("source")),
                       str(entry.get("size", "")), str(entry.get("sha256", "")),
                       "extracted from " + Path(str(entry["derived_from"])).name if entry.get("derived_from") else "original copy"])
    method = [["Step", "Capability", "Depends on", "Required"]]
    steps: list[StepView] = []
    figure_no = 0
    for task in state.plan.steps:
        method.append([f"{task.id}. {task.title}", task.action, ", ".join(map(str, task.depends_on)) or "-",
                       "yes" if task.required else "no"])
        items = [item for item in evidence if item.step_id == task.id]
        figures = []
        for item in items:
            path = (workspace / item.path).resolve()
            if item.type in {"screenshot", "figure"} and integrity.get(item.id, {}).get("verified") and _valid_image(path):
                figure_no += 1
                figures.append(Figure(path, f"{labels['figure']} {figure_no} - {labels['result'].lower()} {task.id}: {item.description}"))
        steps.append(StepView(task.id, task.title, task.status.value, task.action, redact(task.parameters), task.result_summary,
                              task.status_reason, task.requirement_refs, task.verification_results, task.report_sections,
                              figures, items, len(task.attempts), task.required))
    mapping = [["Requirement", "Step(s)", "Status", "Evidence"]]
    refs: dict[str, list[int]] = {}
    for task in state.plan.steps:
        for ref in task.requirement_refs:
            refs.setdefault(ref, []).append(task.id)
    for ref, ids in refs.items():
        statuses = {state.step(i).status.value for i in ids if state.step(i)}  # type: ignore[union-attr]
        ev = [item.id for item in evidence if item.step_id in ids]
        mapping.append([ref[:200], ", ".join(map(str, ids)), "/".join(sorted(statuses)), ", ".join(ev[:8]) or "-"])
    evidence_rows = [["ID", "Type", "Step", "File", "SHA-256", "Intact"]]
    for item in evidence:
        evidence_rows.append([item.id, item.type, str(item.step_id or "-"), item.path, item.sha256 or "-",
                              "yes" if integrity.get(item.id, {}).get("verified") else "NO"])
    errors = [["Step", "Kind", "Message"]] + [[str(e.get("step_id", "-")), str(e.get("kind", "")), str(e.get("error", ""))[:500]]
                                              for e in state.errors]
    recoveries = [["Step", "Attempt", "Recovery"]] + [[str(r.get("step_id")), str(r.get("attempt")), str(r.get("action"))[:300]]
                                                     for r in state.recoveries]
    completed = [t for t in state.plan.steps if t.status == TaskStatus.COMPLETED]
    open_steps = [t for t in state.plan.steps if t.status != TaskStatus.COMPLETED]
    facts = [f"Step {t.id} ({t.title}): {t.result_summary or 'verified'}" for t in completed]
    limitations = [f"Step {t.id} ({t.title}) is {t.status.value}: {t.status_reason or 'not executed'}" for t in open_steps]
    limitations += [f"{a.get('display_name') or n} not available on this computer" for n, a in (environment.get("applications") or {}).items()
                    if not a.get("available") and n in used_tools]
    if any("model" in json.dumps(s.sections, ensure_ascii=False).lower() for s in steps):
        limitations.append("Some checks are model-based (derived from specifications) and are labelled as such.")
    final = (f"{len(completed)} of {len(state.plan.steps)} steps completed and verified; run status {state.status}"
             + (f" ({state.status_reason})." if state.status_reason else "."))
    commands = []
    log = workspace / "logs" / "commands.jsonl"
    if log.is_file():
        for line in log.read_text(encoding="utf-8").splitlines()[-200:]:
            try:
                record = json.loads(line)
            except json.JSONDecodeError:
                continue
            commands.append(f"[{record.get('timestamp', '')}] exit={record.get('exit_code')} :: {' '.join(map(str, record.get('command', [])))}")
    return ReportModel(title=title, assignment=state.assignment, labels=labels, cover=cover, status=str(state.status),
                       status_reason=state.status_reason, objective=state.plan.objective or analysis.get("objective", ""),
                       environment=env_rows, os_line=str(environment.get("os", "unknown OS")), inputs=inputs, method=method,
                       steps=steps, mapping=mapping, evidence=evidence_rows, errors=errors, recoveries=recoveries, facts=facts,
                       limitations=limitations, final=final, commands=commands)


# ---------------------------------------------------------------------------- DOCX
def generate_report(workspace: Path) -> Path:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.image.exceptions import UnrecognizedImageError
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError("DOCX reports require the documents extra: pip install -e .[documents]") from exc
    model = build_model(workspace)
    labels = model.labels
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.8)
    doc.styles["Normal"].font.size = Pt(10.5)

    def table(rows: list[list[str]], widths: list[float] | None = None) -> None:
        if len(rows) <= 1:
            doc.add_paragraph("-")
            return
        grid = doc.add_table(rows=0, cols=len(rows[0]))
        grid.style = "Light Grid Accent 1"
        for index, row in enumerate(rows):
            cells = grid.add_row().cells
            for cell, value in zip(cells, row, strict=False):
                cell.text = str(value)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(8.5)
                        run.bold = index == 0
        if widths:
            for row in grid.rows:
                for cell, width in zip(row.cells, widths, strict=False):
                    cell.width = Inches(width)

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(labels["report"])
    run.font.size, run.bold = Pt(22), True
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(model.title).font.size = Pt(15)
    doc.add_paragraph()
    table([["", ""]] + [[k, v] for k, v in model.cover], [2.0, 4.6])
    doc.add_page_break()

    doc.add_heading(labels["objective"], level=1)
    doc.add_paragraph(model.objective)
    doc.add_heading(labels["environment"], level=1)
    doc.add_paragraph(model.os_line)
    table(model.environment, [1.6, 0.8, 1.0, 3.4])
    doc.add_heading(labels["inputs"], level=1)
    table(model.inputs, [1.8, 0.8, 3.0, 1.2])
    doc.add_heading(labels["method"], level=1)
    table(model.method, [3.4, 2.0, 0.8, 0.6])
    doc.add_heading(labels["execution"], level=1)
    for step in model.steps:
        doc.add_heading(f"{step.id}. {step.title} - {step.status}", level=2)
        if step.requirement_refs:
            doc.add_paragraph("Requirement: " + " | ".join(step.requirement_refs)).runs[0].italic = True
        params = json.dumps(step.parameters, ensure_ascii=False) if step.parameters else "-"
        doc.add_paragraph(f"{labels['action']}: {step.capability} {params}")
        paragraph = doc.add_paragraph(f"{labels['result']}: {step.summary or step.reason or '-'}")
        if step.status != "COMPLETED":
            paragraph.runs[0].font.color.rgb = RGBColor(0xB0, 0x20, 0x20)
        if step.checks:
            doc.add_paragraph(f"{labels['verification']}:")
            table([["Check", "Passed", "Detail"]] + [[c.get("type", ""), "yes" if c.get("passed") else "NO", str(c.get("detail", ""))[:300]]
                                                     for c in step.checks], [1.3, 0.6, 4.9])
        for block in step.sections:
            doc.add_paragraph(block.get("title", "")).runs[0].bold = True
            for text in block.get("paragraphs", []):
                doc.add_paragraph(str(text))
            rows = block.get("table") or []
            if rows:
                keys = list(rows[0].keys())
                table([keys] + [[str(row.get(k, ""))[:120] for k in keys] for row in rows[:60]])
            if block.get("code"):
                code = doc.add_paragraph(str(block["code"])[:6000])
                for run in code.runs:
                    run.font.name, run.font.size = "Consolas", Pt(7.5)
        for figure in step.figures:
            try:
                doc.add_picture(str(figure.path), width=Inches(6.3))
                caption = doc.add_paragraph(figure.caption)
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.runs[0].italic = True
            except (OSError, ValueError, UnrecognizedImageError) as exc:
                doc.add_paragraph(f"Screenshot could not be embedded: {exc}")
        if step.evidence:
            doc.add_paragraph("Evidence: " + "; ".join(f"{e.id} {e.path} (SHA-256 {str(e.sha256)[:16]}...)" for e in step.evidence))
    doc.add_heading(labels["mapping"], level=1)
    table(model.mapping, [3.2, 0.8, 1.2, 1.6])
    doc.add_heading(labels["evidence"], level=1)
    table(model.evidence, [0.7, 0.8, 0.4, 2.2, 2.2, 0.5])
    doc.add_heading(labels["errors"], level=1)
    table(model.errors, [0.5, 1.2, 5.1])
    if len(model.recoveries) > 1:
        table(model.recoveries, [0.5, 0.6, 5.7])
    doc.add_heading(labels["conclusion"], level=1)
    doc.add_heading(labels["facts"], level=2)
    for fact in model.facts or ["-"]:
        doc.add_paragraph(fact, style="List Bullet")
    doc.add_heading(labels["limitations"], level=2)
    for item in model.limitations or ["-"]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading(labels["final"], level=2)
    doc.add_paragraph(model.final)
    doc.add_paragraph("Source evidence remains unchanged; the input manifest records each original path and its hashes.")
    doc.add_heading(labels["appendix"], level=1)
    appendix = doc.add_paragraph("\n".join(model.commands) or "No external commands were executed.")
    for run in appendix.runs:
        run.font.name, run.font.size = "Consolas", Pt(7.5)
    state = _state(workspace)
    docx, _ = _output_paths(workspace, state)
    doc.save(str(docx))
    return docx


# ---------------------------------------------------------------------------- PDF
def generate_pdf_report(workspace: Path) -> Path:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import (
            Image,
            PageBreak,
            Paragraph,
            Preformatted,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
    except ImportError as exc:
        raise RuntimeError("PDF reports require the documents extra: pip install -e .[documents]") from exc
    from xml.sax.saxutils import escape

    model = build_model(workspace)
    labels = model.labels
    styles = getSampleStyleSheet()
    font_candidates = [Path("C:/Windows/Fonts/arial.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
                       Path("/Library/Fonts/Arial Unicode.ttf")]
    mono_candidates = [Path("C:/Windows/Fonts/consola.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf")]
    font, mono = "Helvetica", "Courier"
    regular = next((p for p in font_candidates if p.is_file()), None)
    if regular:
        font = "LabAgentUnicode"
        pdfmetrics.registerFont(TTFont(font, str(regular)))
    monospace = next((p for p in mono_candidates if p.is_file()), None)
    if monospace:
        mono = "LabAgentMono"
        pdfmetrics.registerFont(TTFont(mono, str(monospace)))
    for name in ("Title", "Heading1", "Heading2", "Heading3", "Normal", "BodyText", "Italic"):
        styles[name].fontName = font
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=7.5, leading=9, fontName=font)
    code_style = ParagraphStyle("code", parent=styles["Normal"], fontName=mono, fontSize=6.8, leading=8)
    red = ParagraphStyle("red", parent=styles["Normal"], textColor=colors.HexColor("#B02020"))

    def para(text: Any, style: Any = None) -> Paragraph:
        return Paragraph(escape(str(text)), style or styles["Normal"])

    def table(rows: list[list[str]], widths: list[float] | None = None) -> Any:
        if len(rows) <= 1:
            return para("-")
        data = [[Paragraph(escape(str(cell))[:1500], small) for cell in row] for row in rows]
        grid = Table(data, colWidths=[w * inch for w in widths] if widths else None, repeatRows=1)
        grid.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
                                  ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F2")),
                                  ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        return grid

    story: list[Any] = [Spacer(1, 1.2 * inch), para(labels["report"], styles["Title"]), para(model.title, styles["Heading2"]),
                        Spacer(1, 0.4 * inch), table([["", ""]] + [[k, v] for k, v in model.cover], [2.0, 4.5]), PageBreak()]
    story += [para(labels["objective"], styles["Heading1"]), para(model.objective),
              para(labels["environment"], styles["Heading1"]), para(model.os_line), table(model.environment, [1.5, 0.7, 1.0, 3.5]),
              para(labels["inputs"], styles["Heading1"]), table(model.inputs, [1.7, 0.7, 3.2, 1.2]),
              para(labels["method"], styles["Heading1"]), table(model.method, [3.3, 2.1, 0.7, 0.6]),
              para(labels["execution"], styles["Heading1"])]
    for step in model.steps:
        story.append(para(f"{step.id}. {step.title} - {step.status}", styles["Heading2"]))
        if step.requirement_refs:
            story.append(para("Requirement: " + " | ".join(step.requirement_refs), styles["Italic"]))
        story.append(para(f"{labels['action']}: {step.capability} {json.dumps(step.parameters, ensure_ascii=False) if step.parameters else ''}"))
        story.append(para(f"{labels['result']}: {step.summary or step.reason or '-'}", None if step.status == "COMPLETED" else red))
        if step.checks:
            story.append(table([["Check", "Passed", "Detail"]] + [[c.get("type", ""), "yes" if c.get("passed") else "NO",
                                                                  str(c.get("detail", ""))[:300]] for c in step.checks], [1.3, 0.6, 4.8]))
        for block in step.sections:
            story.append(para(block.get("title", ""), styles["Heading3"]))
            story += [para(text) for text in block.get("paragraphs", [])]
            rows = block.get("table") or []
            if rows:
                keys = list(rows[0].keys())
                story.append(table([keys] + [[str(row.get(k, ""))[:120] for k in keys] for row in rows[:60]]))
            if block.get("code"):
                story.append(Preformatted(str(block["code"])[:6000], code_style, maxLineLength=150))
        for figure in step.figures:
            try:
                picture = Image(str(figure.path))
                picture._restrictSize(6.6 * inch, 7.5 * inch)
                story += [Spacer(1, 0.06 * inch), picture, para(figure.caption, styles["Italic"])]
            except OSError:
                story.append(para("Screenshot could not be embedded."))
        if step.evidence:
            story.append(para("Evidence: " + "; ".join(f"{e.id} {e.path}" for e in step.evidence), small))
        story.append(Spacer(1, 0.12 * inch))
    story += [para(labels["mapping"], styles["Heading1"]), table(model.mapping, [3.2, 0.8, 1.2, 1.5]),
              para(labels["evidence"], styles["Heading1"]), table(model.evidence, [0.6, 0.8, 0.4, 2.2, 2.3, 0.5]),
              para(labels["errors"], styles["Heading1"]), table(model.errors, [0.5, 1.2, 5.0])]
    if len(model.recoveries) > 1:
        story.append(table(model.recoveries, [0.5, 0.6, 5.6]))
    story += [para(labels["conclusion"], styles["Heading1"]), para(labels["facts"], styles["Heading2"])]
    story += [para("• " + fact) for fact in model.facts or ["-"]]
    story += [para(labels["limitations"], styles["Heading2"])] + [para("• " + item) for item in model.limitations or ["-"]]
    story += [para(labels["final"], styles["Heading2"]), para(model.final),
              para(labels["appendix"], styles["Heading1"]),
              Preformatted("\n".join(model.commands) or "No external commands were executed.", code_style, maxLineLength=150)]
    state = _state(workspace)
    _, output = _output_paths(workspace, state)
    SimpleDocTemplate(str(output), pagesize=A4, rightMargin=0.6 * inch, leftMargin=0.6 * inch, topMargin=0.6 * inch,
                      bottomMargin=0.6 * inch, title=f"{model.assignment} {labels['report']}").build(story)
    return output


def generate_reports(workspace: Path, registry: Any = None) -> dict[str, str]:
    """The student report (``<A>_Report``) and the technical audit (``<A>_Audit``), both as DOCX and PDF."""
    from .student_report import generate_student_report

    audit_docx = generate_report(workspace)
    audit_pdf = generate_pdf_report(workspace)
    return {**generate_student_report(workspace, registry), "audit_docx": str(audit_docx), "audit_pdf": str(audit_pdf)}
