from __future__ import annotations

from pathlib import Path

import pytest

from lab_agent.cli import main
from lab_agent.report_builder import (
    ReportSpecError,
    build_report_files,
    build_spec_report,
    lint_file,
)

SPEC = """---
title: Анализ и восстановление разделов MBR и GPT
number: 4
course: Introduction to Digital Forensics
output: Assignment4_{student_id}_{surname}.docx
---
# 1. Цель работы
Проверить целостность образов и восстановить разделы с помощью **TestDisk**.
![Окно TestDisk после анализа](shots/testdisk.png)
Таблица: Контрольные суммы образов
![](data/hashes.csv)

Таблица: Итог
| Образ | Результат |
|---|---|
| mbr.img | восстановлен |
"""


@pytest.fixture
def spec(tmp_path: Path, config) -> Path:  # type: ignore[no-untyped-def]
    from PIL import Image

    config.report.student_name = "Konysbek Abu"
    config.report.group = "cs-2426"
    (tmp_path / "shots").mkdir()
    Image.new("RGB", (800, 500), "navy").save(tmp_path / "shots" / "testdisk.png")
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "hashes.csv").write_text("File,SHA256\nmbr.img,ab12\ngpt.img,cd34\n", encoding="utf-8")
    source = tmp_path / "report.md"
    source.write_text(SPEC, encoding="utf-8")
    return source


def test_markdown_becomes_a_numbered_report(spec: Path) -> None:
    report, _ = build_spec_report(spec)
    assert report.language == "ru" and report.meta["student"] == "Konysbek Abu" and report.meta["number"] == "4"
    captions = [b.text for b in report.blocks if b.kind == "caption"]
    assert captions == ["Рисунок 1 — Окно TestDisk после анализа", "Таблица 1 — Контрольные суммы образов", "Таблица 2 — Итог"]
    tables = [b.rows for b in report.blocks if b.kind == "table"]
    assert tables[0] == [["File", "SHA256"], ["mbr.img", "ab12"], ["gpt.img", "cd34"]]
    # the picture line is not swallowed into the paragraph above it
    assert any(b.kind == "p" and b.text.endswith("**TestDisk**.") for b in report.blocks)


def test_docx_and_pdf_are_written_with_the_requested_name(spec: Path) -> None:
    from docx import Document
    from pypdf import PdfReader

    result = build_report_files(spec)
    docx = Path(result["docx"])
    assert docx.name == "Assignment4_StudentID_Konysbek.docx"
    document = Document(str(docx))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Konysbek Abu" in text and "Рисунок 1 — Окно TestDisk после анализа" in text
    assert len(document.inline_shapes) == 1 and len(document.tables) == 2
    assert len(PdfReader(result["pdf"]).pages) >= 2
    assert result["figures"] == 1 and result["tables"] == 2


def test_missing_picture_is_an_error_not_a_placeholder(spec: Path) -> None:
    spec.write_text(SPEC.replace("shots/testdisk.png", "shots/nope.png"), encoding="utf-8")
    with pytest.raises(ReportSpecError, match="nope.png"):
        build_spec_report(spec)
    (spec.parent / "shots" / "broken.png").write_bytes(b"not a png")
    spec.write_text(SPEC.replace("shots/testdisk.png", "shots/broken.png"), encoding="utf-8")
    with pytest.raises(ReportSpecError, match="not a valid image"):
        build_spec_report(spec)


def test_header_mistakes_are_reported(spec: Path) -> None:
    spec.write_text(SPEC.replace("number: 4", "numbr: 4"), encoding="utf-8")
    with pytest.raises(ReportSpecError, match="numbr"):
        build_spec_report(spec)


def test_lint_flags_machine_phrasing_and_leftovers(spec: Path) -> None:
    spec.write_text(SPEC + "\nВ данной лабораторной работе мы успешно восстановили разделы. [вставить скриншот]\n",
                    encoding="utf-8")
    findings = lint_file(spec)
    assert any("в данной лабораторной работе" in item for item in findings)
    assert any(item.startswith("[error]") and "вставить" in item for item in findings)


def test_lint_is_quiet_on_plain_writing(spec: Path) -> None:
    assert lint_file(spec) == []


def test_cli_build_report_and_strict_mode(spec: Path, capsys: pytest.CaptureFixture[str]) -> None:
    # the CLI loads its own configuration: give it one instead of whatever config.yaml is in the current folder
    settings = spec.parent / "config.yaml"
    settings.write_text("report:\n  student_name: Konysbek Abu\n", encoding="utf-8")
    cli = ["--config", str(settings)]
    assert main([*cli, "build-report", str(spec), "--no-pdf"]) == 0
    assert (spec.parent / "Assignment4_StudentID_Konysbek.docx").is_file()
    spec.write_text(SPEC + "\nTODO дописать вывод\n", encoding="utf-8")
    assert main([*cli, "build-report", str(spec), "--strict"]) == 4
    assert main([*cli, "lint-report", str(spec)]) == 4
    assert main([*cli, "lint-report", str(spec.parent / "Assignment4_StudentID_Konysbek.docx")]) == 0
    capsys.readouterr()
