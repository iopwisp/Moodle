from __future__ import annotations

import csv
import hashlib
import zipfile
from pathlib import Path

import pytest
from assignment_factory import CLUSTER_A, CLUSTER_B, build_assignment_zip, build_image
from fakes import fake_screenshot

from lab_agent.evidence import list_evidence, validate_evidence
from lab_agent.integrations.base import ExecutionContext
from lab_agent.integrations.forensics import detect_signature, parse_scalpel_conf, windows_to_wsl
from lab_agent.runner import run_assignment
from lab_agent.verification import CheckContext, run_checks
from lab_agent.workspace import load_state


def _workspace(tmp_path: Path) -> ExecutionContext:
    image = build_image()
    (tmp_path / "input").mkdir()
    (tmp_path / "input" / "evidence.dd").write_bytes(image)
    (tmp_path / "input" / "evidence.dd.sha256").write_text(hashlib.sha256(image).hexdigest() + "  evidence.dd\n")
    from assignment_factory import _png

    (tmp_path / "input" / "system_log.txt").write_bytes(_png())
    (tmp_path / "input" / "notes.txt").write_text("plain notes", encoding="utf-8")
    return ExecutionContext(tmp_path, "A3", step_id=1, environment={"_wsl": {"available": False, "tools": {}}})


def _checks(result, context: ExecutionContext) -> bool:
    report = run_checks(result.checks, CheckContext(context.workspace, details=result.details,
                                                    evidence_paths=[Path(e["path"]) for e in result.evidence]))
    assert report.passed, report.summary()
    return True


def test_image_hash_working_copy_and_signatures(tmp_path: Path, registry) -> None:
    context = _workspace(tmp_path)
    before = registry.execute("forensics.verify_image_hash", {"label": "before"}, context)
    assert before.verified and before.details["status"] == "SOURCE_UNCHANGED" and _checks(before, context)
    copy = registry.execute("forensics.working_copy", {}, context)
    assert copy.verified and _checks(copy, context)

    scan = registry.execute("forensics.signature_scan", {}, context)
    assert _checks(scan, context)
    rows = {row["file"]: row for row in csv.DictReader((tmp_path / "results" / "file_signature_analysis.csv").open(encoding="utf-8"))}
    assert rows["input/system_log.txt"]["match"] == "no"
    assert rows["input/system_log.txt"]["detected_type"] == "png"
    assert rows["input/system_log.txt"]["signature_hex"].startswith("89 50 4E 47 0D 0A 1A 0A")
    assert rows["input/notes.txt"]["match"] == "yes"

    offsets = registry.execute("forensics.image_signatures", {}, context)
    types = {row["type"] for row in offsets.report_sections[0]["table"]}
    assert {"jpg", "pdf", "png", "zip-eocd"} <= types


def test_tampered_image_is_not_reported_unchanged(tmp_path: Path, registry) -> None:
    context = _workspace(tmp_path)
    image = tmp_path / "input" / "evidence.dd"
    image.write_bytes(image.read_bytes() + b"x")
    result = registry.execute("forensics.verify_image_hash", {"label": "after"}, context)
    assert not result.verified
    assert result.details["status"] == "MISMATCH"


def test_builtin_carver_recovers_valid_files(tmp_path: Path, registry) -> None:
    context = _workspace(tmp_path)
    result = registry.execute("forensics.carve", {"tool": "builtin", "types": "pdf,jpg,png", "output": "builtin_output"}, context)
    assert result.verified and _checks(result, context)
    carved = {row["name"].split("/")[0]: row for row in result.details["files"]}
    assert set(carved) == {"jpg", "pdf", "png"}
    assert all(row["valid"] for row in result.details["files"])
    assert (tmp_path / "results" / "builtin_output" / "audit.txt").read_text().count("valid") >= 3

    hashes = registry.execute("forensics.hash_directory", {}, context)
    assert _checks(hashes, context)
    compare = registry.execute("forensics.compare_results", {}, context)
    assert compare.verified and compare.details["unique_files"] == 3


def test_external_carver_missing_is_blocked_not_faked(tmp_path: Path, registry, monkeypatch) -> None:
    monkeypatch.setattr("shutil.which", lambda name: None)
    context = _workspace(tmp_path)
    result = registry.execute("forensics.carve", {"tool": "foremost", "types": "pdf,jpg"}, context)
    assert result.blocked and not result.verified
    assert not (tmp_path / "results" / "foremost_output").exists()


def test_foremost_runs_through_command_runner(tmp_path: Path, registry, monkeypatch) -> None:
    context = _workspace(tmp_path)
    context.environment["foremost"] = {"available": True, "path": "/usr/bin/foremost"}
    calls = []

    def fake_run(command, **kwargs):
        calls.append(command)
        output = Path(command[command.index("-o") + 1])
        (output / "pdf").mkdir(parents=True)
        (output / "pdf" / "00000024.pdf").write_bytes(b"%PDF-1.4 test")
        (output / "audit.txt").write_text("Foremost version 1.5.7\n1 FILES EXTRACTED\n")
        from lab_agent.tools.process import CommandResult

        return CommandResult(command, 0, "done", "", 0.1)

    monkeypatch.setattr("lab_agent.integrations.forensics.run_command", fake_run)
    result = registry.execute("forensics.carve", {"tool": "foremost", "types": "pdf,jpg", "output": "foremost_output"}, context)
    assert result.verified
    assert calls[0][:2] == ["/usr/bin/foremost", "-v"] and "-t" in calls[0] and "pdf,jpg" in calls[0]
    assert (tmp_path / "results" / "foremost_audit.txt").is_file()


def test_zip_fragment_repair_finds_real_eocd(tmp_path: Path, registry) -> None:
    context = _workspace(tmp_path)
    result = registry.execute("forensics.repair_zip_fragments",
                              {"offsets": f"{hex(CLUSTER_A)},{hex(CLUSTER_B)}", "cluster_size": 4096}, context)
    assert result.verified and _checks(result, context)
    assert result.details["original_header_hex"] == "41 42 43 44"
    with zipfile.ZipFile(tmp_path / "results" / "evidence_fixed.zip") as archive:
        assert archive.read("document.txt") == b"Confidential budget 2026\n"
    script = (tmp_path / "results" / "carve_script_completed.py").read_text()
    assert "PK\\x05\\x06" in script


def test_tsk_missing_is_blocked(tmp_path: Path, registry, monkeypatch) -> None:
    monkeypatch.setattr("shutil.which", lambda name: None)
    result = registry.execute("forensics.tsk", {"command": "fsstat"}, _workspace(tmp_path))
    assert result.blocked


def test_helpers() -> None:
    assert detect_signature(b"\x89PNG\r\n\x1a\n....")[0] == "png"
    assert detect_signature(b"%PDF-1.7")[0] == "pdf"
    rules = parse_scalpel_conf("# jpg y 1 \\xff\\xd8 \\xff\\xd9\npng y 5000000 \\x89\\x50\\x4e\\x47 \\x49\\x45\\x4e\\x44\nzip y 10 PK\\x03\\x04 REVERSE")
    assert [r.extension for r in rules] == ["png", "zip"]
    assert rules[0].header == b"\x89PNG" and rules[1].mode == "REVERSE"
    assert windows_to_wsl(Path("C:/Users/x/image.dd")).startswith("/mnt/c/") or not str(Path("C:/")).startswith("C:")


def test_assignment3_style_pipeline_end_to_end(tmp_path: Path, config, registry) -> None:
    """Assignment ZIP -> analyzer -> planner -> runner -> verification -> evidence -> DOCX/PDF."""
    upload = build_assignment_zip(tmp_path)
    workspace, state, planner = run_assignment([upload], tmp_path / "workspaces", "deterministic", config=config,
                                               registry=registry, screenshot_fn=fake_screenshot,
                                               environment={"os": "test", "applications": {}, "wsl": {"available": False, "tools": {}}})
    by_action = {}
    for task in state.plan.steps:
        by_action.setdefault(task.action, []).append(task)
    assert planner == "deterministic"
    # Real, verifiable forensic work completed:
    for action in ("forensics.verify_image_hash", "forensics.working_copy", "forensics.signature_scan",
                   "forensics.image_signatures", "forensics.hash_directory", "forensics.compare_results",
                   "forensics.repair_zip_fragments", "forensics.case_records"):
        assert all(t.status.value == "COMPLETED" for t in by_action[action]), (action, [t.status_reason for t in by_action[action]])
    builtin = next(t for t in by_action["forensics.carve"] if t.parameters["tool"] == "builtin")
    assert builtin.status.value == "COMPLETED"
    # Honest failure behaviour: tools that are not installed are BLOCKED, never COMPLETED.
    foremost = next(t for t in by_action["forensics.carve"] if t.parameters["tool"] == "foremost")
    assert foremost.status.value == "BLOCKED"
    assert by_action["autopsy.ingest"][0].status.value == "BLOCKED"
    assert by_action["core.manual_review"][0].status.value == "BLOCKED"
    assert state.status in {"BLOCKED", "WAITING_CONFIRMATION"}
    # the image inside the upload was extracted, hashed before/after and never modified
    assert "SOURCE_UNCHANGED" in (workspace / "results" / "source_image_sha256_after.txt").read_text()
    evidence = list_evidence(workspace)
    assert evidence and all(item["verified"] for item in validate_evidence(workspace))
    assert any(item.path.endswith("evidence_fixed.zip") for item in evidence)
    for deliverable in ("file_signature_analysis.csv", "recovered_files_sha256.csv", "comparison_results.csv",
                        "chain_of_custody.csv", "case_context.csv", "commands.txt"):
        assert (workspace / "results" / deliverable).is_file(), deliverable
    report_docx = workspace / "reports" / f"{state.assignment}_Report.docx"
    report_pdf = workspace / "reports" / f"{state.assignment}_Report.pdf"
    assert report_docx.is_file() and report_pdf.is_file()
    from docx import Document

    text = "\n".join(p.text for p in Document(str(report_docx)).paragraphs)
    assert "BLOCKED" in text and "SOURCE_UNCHANGED" in "\n".join(
        c.text for t in Document(str(report_docx)).tables for r in t.rows for c in r.cells) + text
    assert load_state(workspace).status == state.status


@pytest.mark.skipif(not (Path.home() / "Downloads" / "Assignment_3_Student_Materials.zip").is_file(),
                    reason="real Assignment 3 materials are only available on the author's machine")
def test_real_assignment3_materials_local(tmp_path: Path, config, registry) -> None:
    upload = Path.home() / "Downloads" / "Assignment_3_Student_Materials.zip"
    workspace, state, _ = run_assignment([upload], tmp_path / "ws", "deterministic", config=config, registry=registry,
                                         screenshot_fn=fake_screenshot, environment={"os": "local", "applications": {}, "wsl": {"available": False, "tools": {}}})
    repair = next(t for t in state.plan.steps if t.action == "forensics.repair_zip_fragments")
    assert repair.status.value == "COMPLETED", repair.status_reason
    with zipfile.ZipFile(workspace / "results" / "evidence_fixed.zip") as archive:
        assert archive.namelist() == ["document.txt"]
    assert "SOURCE_UNCHANGED" in (workspace / "results" / "source_image_sha256_before.txt").read_text()
