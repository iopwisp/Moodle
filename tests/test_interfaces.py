"""Reports, web dashboard, CLI, environment discovery and the PowerShell allowlist."""

from __future__ import annotations

import json
import os
from pathlib import Path

import pytest
from conftest import EMPTY_ENVIRONMENT
from fakes import fake_screenshot

from lab_agent.cli import main
from lab_agent.environment import discover_app, format_doctor, run_doctor
from lab_agent.integrations.base import AppSpec, ExecutionContext
from lab_agent.runner import run_assignment
from lab_agent.workspace import load_state


def _run_small(tmp_path: Path, config, registry, text: str = "1. Calculate SHA256 hash\n2. Take a screenshot of the result\n"):
    source = tmp_path / "lab.txt"
    source.write_text(text, encoding="utf-8")
    return run_assignment([source], tmp_path / "ws", "deterministic", config=config, registry=registry,
                          screenshot_fn=fake_screenshot, environment=EMPTY_ENVIRONMENT)


def test_report_is_submission_oriented_and_honest(tmp_path: Path, config, registry) -> None:
    config.report.student_name = "Test Student"
    config.report.group = "CS-2435"
    workspace, state, _ = _run_small(tmp_path, config, registry, "1. Calculate SHA256 hash\n2. Take a screenshot\n3. Write an essay\n")
    from docx import Document
    from pypdf import PdfReader

    # Technical audit: every status, check and the requirement mapping.
    audit = Document(str(workspace / "reports" / f"{state.assignment}_Audit.docx"))
    audit_text = "\n".join(p.text for p in audit.paragraphs)
    for heading in ("1. Objective", "2. Environment", "3. Input materials", "5. Execution and results", "6. Requirement mapping",
                    "7. Evidence register", "9. Conclusion", "Limitations", "Appendix A"):
        assert heading in audit_text, heading
    assert "3. Write an essay - BLOCKED" in audit_text and "Step 3 (Write an essay) is BLOCKED" in audit_text
    audit_pdf = "\n".join(page.extract_text() for page in PdfReader(str(workspace / "reports" / f"{state.assignment}_Audit.pdf")).pages)
    assert "Requirement mapping" in audit_pdf and "BLOCKED" in audit_pdf

    # Student report: title page, plain prose in the assignment's language, honest about what is missing.
    document = Document(str(workspace / "reports" / f"{state.assignment}_Report.docx"))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "LABORATORY REPORT" in text and "Student: Test Student, group CS-2435" in text
    for heading in ("Procedure and results", "Conclusion", "Appendix B. Commands"):
        assert heading in text, heading
    for jargon in ("BLOCKED", "COMPLETED", "core.", "checks passed"):
        assert jargon not in text, jargon
    assert "This step could not be carried out: it cannot be automated and has to be done by hand." in text
    assert "Not completed: “Write an essay”." in text
    assert len(document.inline_shapes) >= 1  # the screenshot of step 2 is embedded
    pdf_text = "\n".join(page.extract_text() for page in PdfReader(str(workspace / "reports" / f"{state.assignment}_Report.pdf")).pages)
    assert "Procedure and results" in pdf_text and "BLOCKED" not in pdf_text


def test_web_dashboard_status_events_controls_and_files(tmp_path: Path, config, registry) -> None:
    from fastapi.testclient import TestClient

    from lab_agent.web import create_app

    root = tmp_path / "ws"
    app = create_app(root, run_options={"config": config, "registry": registry, "screenshot_fn": fake_screenshot,
                                        "environment": EMPTY_ENVIRONMENT})
    client = TestClient(app)
    assert "Lab Agent" in client.get("/").text
    response = client.post("/assignments/analyze", files=[("files", ("lab.txt", b"1. Calculate SHA256 hash\n2. Take a screenshot\n"))],
                           data={"provider": "deterministic"})
    assignment = response.json()["id"]
    app.state.manager.wait(root / assignment, 60)
    status = client.get(f"/assignments/{assignment}/status").json()
    assert status["status"] == "COMPLETED" and status["progress"] == 100 and status["reports"]["pdf"]
    stream = client.get(f"/assignments/{assignment}/events/stream?once=true")
    assert "task.completed" in stream.text and stream.headers["content-type"].startswith("text/event-stream")
    shot = next(e for e in status["evidence"] if e["type"] == "screenshot")
    assert client.get(f"/assignments/{assignment}/files/{shot['path']}").status_code == 200
    assert client.get(f"/assignments/{assignment}/files/state/state.json").status_code == 404
    assert client.get(f"/assignments/{assignment}/files/../../etc/passwd").status_code == 404
    assert client.get("/assignments/../outside/status").status_code == 404
    assert client.get(f"/reports/{assignment}/docx").status_code == 200
    assert client.post(f"/assignments/{assignment}/pause").json()["status"] == "PAUSE_REQUESTED"
    assert client.post(f"/assignments/{assignment}/tasks/1/approve").status_code == 404
    page = client.get(f"/assignments/{assignment}").text
    assert "EventSource" in page and "Pause" in page
    assert any(c["name"] == "forensics.carve" for c in client.get("/api/capabilities").json())


def test_cli_commands(tmp_path: Path, capsys, monkeypatch) -> None:
    monkeypatch.chdir(tmp_path)
    assert main(["capabilities", "--tool", "core"]) == 0
    assert "core.hash_inputs" in capsys.readouterr().out
    assert main(["plugins"]) == 0
    plugins = json.loads(capsys.readouterr().out)
    assert {"core", "forensics", "packet_tracer"} <= {p["name"] for p in plugins["integrations"]}
    assert main(["config"]) == 0 and '"workspace_root"' in capsys.readouterr().out
    source = tmp_path / "lab.txt"
    source.write_text("1. Calculate SHA256 hash\n", encoding="utf-8")
    monkeypatch.setattr("lab_agent.runner.discover_environment", lambda *a, **k: EMPTY_ENVIRONMENT)
    assert main(["plan", str(source), "--workspace-root", str(tmp_path / "ws"), "--ai-provider", "deterministic"]) == 0
    planned = json.loads(capsys.readouterr().out)
    assert planned["steps"][0]["capability"] == "core.hash_inputs"
    workspace = planned["workspace"]
    assert main(["status", workspace]) == 0 and '"PLANNED"' in capsys.readouterr().out
    assert main(["stop", workspace]) == 0 and load_state(Path(workspace)).status == "STOPPED"
    capsys.readouterr()
    assert main(["resume", workspace, "--verify-only"]) == 0
    assert main(["status", str(tmp_path)]) == 2  # not a workspace


def test_doctor_reports_structure(tmp_path: Path, config, registry, monkeypatch) -> None:
    monkeypatch.setattr("lab_agent.environment._ai_status", lambda config: [("OpenAI", "-", "not set")])
    report = run_doctor(registry, config, tmp_path, include_wsl=False)
    groups = {c["group"] for c in report["checks"]}
    assert {"runtime", "dependencies", "applications", "ai", "desktop", "workspace", "plugins"} <= groups
    assert report["result"] in {"READY", "READY_WITH_WARNINGS", "NOT_READY"}
    assert "Result:" in format_doctor(report)


def test_environment_discovery_from_env_and_install_dirs(tmp_path: Path, monkeypatch) -> None:
    exe = tmp_path / "Cisco Packet Tracer 8.2.2" / "bin" / "PacketTracer.exe"
    exe.parent.mkdir(parents=True)
    exe.write_bytes(b"MZ")
    spec = AppSpec("packet_tracer", "Cisco Packet Tracer", env_var="LAB_AGENT_PACKET_TRACER_PATH",
                   executables=("PacketTracer.exe",), install_globs=("Cisco Packet Tracer*/bin/PacketTracer.exe",))
    monkeypatch.setenv("LAB_AGENT_PACKET_TRACER_PATH", str(exe))
    info = discover_app(spec, registry_entries=[])
    assert info.available and info.source.startswith("env:") and info.version == "8.2.2"
    monkeypatch.delenv("LAB_AGENT_PACKET_TRACER_PATH")
    monkeypatch.setattr("lab_agent.environment._roots", lambda: [tmp_path])
    info = discover_app(spec, registry_entries=[])
    assert info.available and info.source.startswith("install_dir")
    missing = discover_app(AppSpec("nothing", "Nothing", env_var="LAB_AGENT_NOTHING"), registry_entries=[])
    assert not missing.available and "LAB_AGENT_NOTHING" in missing.note


@pytest.mark.skipif(os.name != "nt", reason="Windows PowerShell")
def test_powershell_allowlisted_get_filehash(tmp_path: Path, registry) -> None:
    (tmp_path / "input").mkdir()
    (tmp_path / "input" / "e.bin").write_bytes(b"abc")
    result = registry.execute("powershell.get_file_hash", {"path": "input/e.bin"}, ExecutionContext(tmp_path, "a"))
    assert result.verified and result.details["hash"] == result.details["expected"]


def test_arbitrary_powershell_cannot_be_enabled_by_the_plan(tmp_path: Path, config, registry) -> None:
    from lab_agent.policy import PolicyEngine

    capability = registry.capability("powershell.run_script")
    decision = PolicyEngine(config).check(capability, {"command": "Get-Date", "allow_unsafe": True})
    assert not decision.allowed
    assert registry.validate_parameters("powershell.run_script", {"command": "x", "allow_unsafe": True})
